"""Resume text extraction (PDF/DOCX) + lightweight heuristic fields."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Optional

import pdfplumber
from docx import Document


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract raw text from a resume PDF or DOCX upload."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return _pdf_text(file_bytes)
    if name.endswith(".docx"):
        return _docx_text(file_bytes)
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume format: {filename}")


def _pdf_text(data: bytes) -> str:
    out: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            txt = page.extract_text() or ""
            out.append(txt)
    return "\n".join(out).strip()


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paras.extend(cell.text for cell in row.cells)
    return "\n".join([p for p in paras if p.strip()])


# --- Heuristic fallbacks used before AI runs / for sanity checks -------

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_YEAR_RE = re.compile(r"(\d+)\+?\s+years?", re.I)
_COMMON_SKILLS = {
    "python", "sql", "java", "javascript", "typescript", "go", "rust", "c++",
    "react", "node", "django", "flask", "fastapi", "spring", "kubernetes",
    "docker", "aws", "gcp", "azure", "terraform", "snowflake", "bigquery",
    "airflow", "dbt", "spark", "kafka", "pytorch", "tensorflow", "pandas",
    "numpy", "scikit-learn", "huggingface", "llm", "nlp", "ml", "etl",
    "postgresql", "mysql", "mongodb", "redis", "graphql", "rest", "git",
}


def quick_fields(text: str) -> dict:
    """Run regex/heuristic field extraction."""
    lower = text.lower()
    found_skills = sorted({s for s in _COMMON_SKILLS if re.search(rf"\b{re.escape(s)}\b", lower)})
    years_match = _YEAR_RE.search(lower)
    return {
        "email": (_EMAIL_RE.search(text) or [None])[0] if _EMAIL_RE.search(text) else None,
        "skills_guess": found_skills,
        "years_experience_guess": float(years_match.group(1)) if years_match else None,
    }


def save_upload(file_bytes: bytes, filename: str, folder: str = "resumes") -> Path:
    """Persist uploaded resume to disk and return path."""
    Path(folder).mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", filename)
    path = Path(folder) / safe
    path.write_bytes(file_bytes)
    return path
